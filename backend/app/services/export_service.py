"""
Export Service for Contact Data
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from typing import Optional
from datetime import datetime
from pathlib import Path
import logging
from io import BytesIO

from app.models.contact import Contact
from app.config import settings

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting contact data to Excel"""
    
    def __init__(self, db: Session):
        self.db = db
        self.export_dir = Path(settings.EXPORT_DIR)
        self.export_dir.mkdir(parents=True, exist_ok=True)
    
    def export_contacts_to_excel(
        self,
        filename: str,
        country: Optional[str] = None,
        keyword: Optional[str] = None,
        city: Optional[str] = None,
        institution_type: Optional[str] = None,
        source_platform: Optional[str] = None,
        min_quality_score: Optional[float] = None,
        has_email: Optional[bool] = None,
        has_whatsapp: Optional[bool] = None,
        date_from: Optional[datetime] = None,
        date_to: Optional[datetime] = None,
        max_records: Optional[int] = None
    ) -> str:
        """
        Export contacts to Excel file with bilingual headers
        
        Args:
            filename: Output filename
            country: Filter by country code
            keyword: Filter by search keyword
            city: Filter by city (requires join with Task table)
            institution_type: Filter by institution type
            source_platform: Filter by source platform
            min_quality_score: Minimum quality score filter
            has_email: Filter by email presence
            has_whatsapp: Filter by WhatsApp presence
            date_from: Filter by extraction date (from)
            date_to: Filter by extraction date (to)
            max_records: Maximum number of records to export
            
        Returns:
            Full path to the exported file
        """
        logger.info(f"Starting export to {filename}")
        
        df = self._build_contacts_dataframe(
            country=country,
            keyword=keyword,
            city=city,
            institution_type=institution_type,
            source_platform=source_platform,
            min_quality_score=min_quality_score,
            has_email=has_email,
            has_whatsapp=has_whatsapp,
            date_from=date_from,
            date_to=date_to,
            max_records=max_records
        )
        
        filepath = self.export_dir / filename
        workbook = self._create_workbook(df)
        workbook.save(filepath)
        workbook.close()
        
        records_count = len(df.index)
        file_size = filepath.stat().st_size
        
        logger.info(f"Export completed: {filepath}")
        return str(filepath)

    def export_contacts_to_excel_bytes(
        self,
        filename: str,
        country: Optional[str] = None,
        keyword: Optional[str] = None,
        city: Optional[str] = None,
        institution_type: Optional[str] = None,
        source_platform: Optional[str] = None,
        min_quality_score: Optional[float] = None,
        has_email: Optional[bool] = None,
        has_whatsapp: Optional[bool] = None,
        date_from: Optional[datetime] = None,
        date_to: Optional[datetime] = None,
        max_records: Optional[int] = None
    ) -> dict:
        """Export contacts to Excel and return bytes content."""
        df = self._build_contacts_dataframe(
            country=country,
            keyword=keyword,
            city=city,
            institution_type=institution_type,
            source_platform=source_platform,
            min_quality_score=min_quality_score,
            has_email=has_email,
            has_whatsapp=has_whatsapp,
            date_from=date_from,
            date_to=date_to,
            max_records=max_records
        )

        workbook = self._create_workbook(df)
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        buffer.seek(0)

        content = buffer.getvalue()
        logger.info(f"Export generated in-memory: {filename} ({len(df.index)} records, {len(content)} bytes)")

        return {
            "filename": filename,
            "content": content,
            "records_count": len(df.index),
            "file_size": len(content)
        }

    def _build_contacts_dataframe(
        self,
        country: Optional[str] = None,
        keyword: Optional[str] = None,
        city: Optional[str] = None,
        institution_type: Optional[str] = None,
        source_platform: Optional[str] = None,
        min_quality_score: Optional[float] = None,
        has_email: Optional[bool] = None,
        has_whatsapp: Optional[bool] = None,
        date_from: Optional[datetime] = None,
        date_to: Optional[datetime] = None,
        max_records: Optional[int] = None
    ) -> pd.DataFrame:
        from app.models.task import Task

        if city:
            query = self.db.query(Contact).join(Task, Contact.task_id == Task.id)
        else:
            query = self.db.query(Contact)

        if country:
            query = query.filter(Contact.country == country)

        if keyword:
            query = query.filter(Contact.keyword == keyword)

        if city:
            query = query.filter(Task.city == city)

        if institution_type:
            query = query.filter(Contact.institution_type == institution_type)

        if source_platform:
            query = query.filter(Contact.source_platform == source_platform)

        if min_quality_score is not None:
            query = query.filter(Contact.quality_score >= min_quality_score)

        if has_email is not None:
            if has_email:
                query = query.filter(Contact.email.isnot(None))
            else:
                query = query.filter(Contact.email.is_(None))

        if has_whatsapp is not None:
            if has_whatsapp:
                query = query.filter(Contact.whatsapp.isnot(None))
            else:
                query = query.filter(Contact.whatsapp.is_(None))

        if date_from:
            query = query.filter(Contact.extracted_at >= date_from)

        if date_to:
            query = query.filter(Contact.extracted_at <= date_to)

        query = query.order_by(Contact.extracted_at.desc())

        if max_records:
            query = query.limit(max_records)
        else:
            query = query.limit(settings.EXPORT_MAX_RECORDS)

        contacts = query.all()

        if not contacts:
            logger.warning("No contacts found matching the filters")
            raise ValueError("No contacts found matching the filters")

        logger.info(f"Preparing export for {len(contacts)} contacts")

        data = []
        for contact in contacts:
            data.append({
                'id': str(contact.id),
                'task_id': str(contact.task_id),
                'country': contact.country,
                'institution_name': contact.institution_name,
                'institution_type': contact.institution_type or '',
                'source_url': contact.source_url,
                'source_platform': contact.source_platform,
                'email': contact.email or '',
                'whatsapp': contact.whatsapp or '',
                'quality_score': contact.quality_score or 0,
                'is_verified': 'Yes' if contact.is_verified else 'No',
                'extracted_at': contact.extracted_at.strftime('%Y-%m-%d %H:%M:%S') if contact.extracted_at else ''
            })

        return pd.DataFrame(data)
    
    def _create_workbook(self, df: pd.DataFrame) -> Workbook:
        wb = Workbook()
        ws = wb.active
        ws.title = "Contacts"
        
        # Define bilingual headers (Chinese / English)
        headers = [
            ('ID', 'ID'),
            ('任務ID', 'Task ID'),
            ('國家', 'Country'),
            ('機構名稱', 'Institution Name'),
            ('機構類型', 'Institution Type'),
            ('來源網址', 'Source URL'),
            ('來源平台', 'Source Platform'),
            ('Email', 'Email'),
            ('WhatsApp', 'WhatsApp'),
            ('品質分數', 'Quality Score'),
            ('已驗證', 'Verified'),
            ('萃取時間', 'Extracted At')
        ]
        
        # Header styles
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Write Chinese headers (row 1)
        for col_idx, (chinese, english) in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = chinese
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        
        # Write English headers (row 2)
        for col_idx, (chinese, english) in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col_idx)
            cell.value = english
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        
        # Write data rows
        for row_idx, row_data in enumerate(df.values, start=3):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=False)
        
        # Adjust column widths
        column_widths = {
            'A': 38,  # ID
            'B': 38,  # Task ID
            'C': 10,  # Country
            'D': 40,  # Institution Name
            'E': 15,  # Institution Type
            'F': 50,  # Source URL
            'G': 15,  # Source Platform
            'H': 30,  # Email
            'I': 20,  # WhatsApp
            'J': 12,  # Quality Score
            'K': 10,  # Verified
            'L': 20   # Extracted At
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # Set row heights
        ws.row_dimensions[1].height = 25
        ws.row_dimensions[2].height = 25
        
        # Freeze header rows
        ws.freeze_panes = 'A3'
        
        return wb
